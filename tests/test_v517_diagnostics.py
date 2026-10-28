"""Regression tests for stack checks and downloadable Diagnosis reports."""
import csv
import io
import unittest

import ios_xe_upgrade_manager as manager


class StackSnapshotTests(unittest.TestCase):
    def test_stack_ports_parser_accepts_member_status_rows(self):
        self.assertEqual(
            manager.parse_stack_ports("Switch Port1 Port2\n1 OK OK\n2 DOWN OK\n"),
            {"1": ["OK", "OK"], "2": ["DOWN", "OK"]},
        )

    def test_stack_ports_parser_fails_closed_for_cli_errors(self):
        self.assertEqual(manager.parse_stack_ports("% Invalid input detected"), {})

    def test_stack_power_parser_extracts_stack_summary(self):
        parsed = manager.parse_stack_power("PowerStack1 SP-PS Redundant 3000 3 3 4 2 4\n")
        self.assertEqual(parsed, {
            "PowerStack1": {"mode": "SP-PS", "topology": "Redundant", "total": 3000, "sw": 2, "ps": 4}
        })

    def test_compare_snapshots_reports_stack_member_port_and_power_changes(self):
        pre = {
            "ok": {"stack_members": True, "stack_ports": True, "stack_power": True},
            "stack_members": {"1": {"role": "Active", "state": "Ready"}, "2": {"role": "Member", "state": "Ready"}},
            "stack_ports": {"1": ["OK", "OK"], "2": ["OK", "OK"]},
            "stack_power": {"PowerStack1": {"topology": "Redundant", "sw": 2, "ps": 4}},
        }
        post = {
            "ok": {"stack_members": True, "stack_ports": True, "stack_power": True},
            "stack_members": {"1": {"role": "Active", "state": "Provisioned"}},
            "stack_ports": {"1": ["DOWN", "OK"]},
            "stack_power": {"PowerStack1": {"topology": "Combined", "sw": 1, "ps": 2}},
        }
        issues = manager.compare_snapshots(pre, post)
        self.assertTrue(any("switch 2" in issue and "missing" in issue for issue in issues))
        self.assertTrue(any("no longer Ready" in issue for issue in issues))
        self.assertTrue(any("stack port(s)" in issue for issue in issues))
        self.assertTrue(any("topology" in issue for issue in issues))
        self.assertTrue(any("power supply" in issue for issue in issues))


class DiagnosisReportTests(unittest.TestCase):
    @staticmethod
    def analysis():
        return {
            "label": "Switch A", "host": "192.0.2.10", "verdict": "amber",
            "verdict_text": "Review one warning", "time": "2026-10-28 10:00",
            "lookback_hours": 24, "findings": [{
                "level": "warn", "category": "test", "title": "=1+1", "count": 1,
                "first": "10:00", "last": "10:00", "advice": "Review the switch",
                "examples": ["sample evidence"],
            }],
            "buffer": {"lines": 10, "first": "09:00", "last": "10:00"},
            "facts": {"version": "17.12.4"}, "notes": [], "files": {}, "top": [],
        }

    def test_csv_export_has_bom_and_neutralizes_formula_cells(self):
        content = manager.build_diag_csv([self.analysis()])
        self.assertTrue(content.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        self.assertIn("'=1+1", rows[1])

    def test_excel_export_contains_summary_findings_and_safe_text(self):
        from openpyxl import load_workbook

        content = manager.build_diag_xlsx([self.analysis()])
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Summary", "Findings", "Top messages"])
        self.assertEqual(workbook["Findings"]["H2"].value, "=1+1")
        self.assertEqual(workbook["Findings"]["H2"].data_type, "s")

    def test_pdf_export_has_pdf_signature_and_end_marker(self):
        content = manager.build_diag_pdf([self.analysis()], site="Lab")
        self.assertTrue(content.startswith(b"%PDF-1.4"))
        self.assertTrue(content.rstrip().endswith(b"%%EOF"))

    def test_report_dispatch_rejects_unsupported_formats(self):
        with self.assertRaises(ValueError):
            manager.build_diag_report("html", [self.analysis()])


if __name__ == "__main__":
    unittest.main()
